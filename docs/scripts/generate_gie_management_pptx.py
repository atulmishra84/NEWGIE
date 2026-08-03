#!/usr/bin/env python3
"""Generate management-review PPTX: GIE as guardrail/policy orchestration layer."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

INK = RGBColor(0x0A, 0x16, 0x28)
NAVY = RGBColor(0x10, 0x2A, 0x43)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
TEAL_BRIGHT = RGBColor(0x14, 0xB8, 0xA6)
MIST = RGBColor(0x9F, 0xB3, 0xC8)
FOG = RGBColor(0xD9, 0xE2, 0xEC)
PAPER = RGBColor(0xF0, 0xF4, 0xF8)
WHITE = RGBColor(0xFA, 0xFB, 0xFC)
SLATE = RGBColor(0x48, 0x65, 0x81)

OUT = Path(__file__).resolve().parents[1] / "presentations" / "GIE_Management_Orchestration_Layer.pptx"


def strip_style(shape) -> None:
    sp = shape._element
    for style in sp.findall(qn("p:style")):
        sp.remove(style)


def set_solid(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    strip_style(shape)


def add_rect(slide, left, top, width, height, color: RGBColor):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    set_solid(sh, color)
    return sh


def textbox(slide, left, top, width, height, lines, *, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Calibri"
        p.space_after = Pt(6)
    return box


def blank(prs):
    return prs.slide_layouts[6]


def title_bar(slide, title: str) -> None:
    add_rect(slide, 0, 0, 13.333, 0.95, NAVY)
    add_rect(slide, 0, 0.95, 13.333, 0.08, TEAL_BRIGHT)
    textbox(slide, 0.5, 0.22, 12.2, 0.6, [title], size=28, bold=True, color=WHITE)


def bullets(slide, left, top, width, height, items, *, size=18, color=INK):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        run = p.add_run()
        run.text = f"•  {item}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Calibri"
        p.space_after = Pt(8)
    return box


def card(slide, left, top, width, height, title, body, *, accent=False):
    bg = WHITE
    sh = add_rect(slide, left, top, width, height, bg)
    if accent:
        sh.line.color.rgb = TEAL_BRIGHT
        sh.line.width = Pt(1.5)
    else:
        sh.line.color.rgb = FOG
        sh.line.width = Pt(1)
    textbox(slide, left + 0.15, top + 0.12, width - 0.3, 0.4, [title], size=16, bold=True, color=TEAL)
    textbox(slide, left + 0.15, top + 0.5, width - 0.3, height - 0.65, body, size=13, color=SLATE)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 Title
    s = prs.slides.add_slide(blank(prs))
    add_rect(s, 0, 0, 13.333, 7.5, NAVY)
    add_rect(s, 0, 5.9, 13.333, 1.6, RGBColor(0x07, 0x13, 0x1F))
    textbox(s, 0.6, 2.2, 12, 1.2, ["GIE"], size=60, bold=True, color=TEAL_BRIGHT)
    textbox(s, 0.6, 3.3, 12, 0.4, ["GUARDRAILS INTELLIGENCE ENGINE"], size=14, bold=True, color=MIST)
    textbox(
        s,
        0.6,
        4.0,
        11,
        1.2,
        ["Orchestration layer for choosing the right", "guardrails & policies — Management Review"],
        size=24,
        bold=True,
        color=WHITE,
    )
    textbox(
        s,
        0.6,
        6.2,
        12,
        0.8,
        ["How GIE helps the agentic enterprise select, generate, and gate controls before AI agents ship."],
        size=14,
        color=MIST,
    )

    # 2 Agenda
    s = prs.slides.add_slide(blank(prs))
    add_rect(s, 0, 0, 13.333, 7.5, PAPER)
    title_bar(s, "Agenda")
    bullets(
        s,
        0.7,
        1.4,
        11.5,
        5,
        [
            "Problem in today’s agentic world",
            "What GIE is (management view)",
            "How it chooses the right controls",
            "Where to leverage it (layers) — key discussion",
            "Business outcomes & USP",
            "What it is / is not",
            "Adoption path & ask of this review",
        ],
        size=20,
    )

    # 3 Problem
    s = prs.slides.add_slide(blank(prs))
    add_rect(s, 0, 0, 13.333, 7.5, PAPER)
    title_bar(s, "The management problem")
    card(s, 0.5, 1.4, 3.9, 3.2, "Agents are proliferating", [
        "Teams ship autonomous agents with tools, MCP, memory, and multi-agent graphs faster than policy teams can keep up."
    ])
    card(s, 4.7, 1.4, 3.9, 3.2, "One-size policies fail", [
        "A chatbot and a tool-using ops agent need different guardrails. Generic packs over-block or under-protect."
    ])
    card(s, 8.9, 1.4, 3.9, 3.2, "No release decision", [
        "Organizations lack a standard answer: “Is this agent clear to ship — and with which controls?”"
    ])
    add_rect(s, 0.5, 5.0, 12.3, 1.5, NAVY)
    textbox(
        s,
        0.75,
        5.35,
        11.8,
        1.0,
        ["Without an orchestration layer, policy selection stays manual, inconsistent, and hard to audit."],
        size=16,
        bold=True,
        color=WHITE,
    )

    # 4 What GIE is
    s = prs.slides.add_slide(blank(prs))
    add_rect(s, 0, 0, 13.333, 7.5, PAPER)
    title_bar(s, "What GIE is")
    textbox(
        s,
        0.6,
        1.4,
        12,
        1.4,
        [
            "GIE is the governance orchestration layer that helps the organization",
            "choose the right guardrails and policies for every AI agent.",
        ],
        size=22,
        bold=True,
        color=NAVY,
    )
    card(
        s,
        0.5,
        3.2,
        6.0,
        2.8,
        "For leadership",
        ["One control plane: discover → risk → recommend → generate policies → validate → explain → release decision."],
        accent=True,
    )
    card(
        s,
        6.8,
        3.2,
        6.0,
        2.8,
        "For builders & security",
        ["Context-aware controls based on what the agent actually is and what it can do — not questionnaires alone."],
    )

    # 5 How it chooses
    s = prs.slides.add_slide(blank(prs))
    add_rect(s, 0, 0, 13.333, 7.5, PAPER)
    title_bar(s, "How GIE chooses the right guardrails")
    steps = [
        ("1. Discover", "Inventory: frameworks, models, tools, MCP, data signals"),
        ("2. Assess", "Risk + compliance gaps for that agent"),
        ("3. Recommend", "Prioritized guardrails matched to findings"),
        ("4. Generate", "Deployment-ready policy artifacts"),
        ("5. Validate", "Pre-deploy checks + release posture"),
    ]
    x = 0.4
    for title, body in steps:
        card(s, x, 1.5, 2.35, 3.4, title, [body])
        x += 2.55
    add_rect(s, 0.5, 5.3, 12.3, 1.4, NAVY)
    textbox(
        s,
        0.75,
        5.65,
        11.8,
        0.9,
        ["Output: right controls for this agent + audit ID + hold / conditional / clear decision."],
        size=16,
        bold=True,
        color=WHITE,
    )

    # 6 Layers table
    s = prs.slides.add_slide(blank(prs))
    add_rect(s, 0, 0, 13.333, 7.5, PAPER)
    title_bar(s, "Where to leverage GIE — enterprise layers")
    rows = [
        ("Build (IDE / repo)", "Developers create agents", "Early discovery + policy hints while coding"),
        ("CI / CD pipeline", "PR / build / release gates", "PRIMARY: analyze every agent change; block unsafe releases"),
        ("Pre-prod governance", "Security / compliance review", "Risk brief, gaps, recommended + generated policies"),
        ("Runtime / ops", "Agent serving, tools, MCP", "Feeds the right firewall / DLP / policy packs to enforce"),
        ("GRC / audit", "Frameworks, evidence, boards", "Evidence pack, audit IDs, explainability"),
    ]
    y = 1.25
    add_rect(s, 0.4, y, 12.5, 0.45, NAVY)
    textbox(s, 0.55, y + 0.05, 2.8, 0.35, ["Layer"], size=12, bold=True, color=WHITE)
    textbox(s, 3.5, y + 0.05, 4.0, 0.35, ["What happens"], size=12, bold=True, color=WHITE)
    textbox(s, 7.7, y + 0.05, 5.0, 0.35, ["How GIE is leveraged"], size=12, bold=True, color=WHITE)
    y = 1.7
    for i, (a, b, c) in enumerate(rows):
        bg = WHITE if i % 2 == 0 else RGBColor(0xE8, 0xEE, 0xF4)
        add_rect(s, 0.4, y, 12.5, 0.75, bg)
        textbox(s, 0.55, y + 0.15, 2.8, 0.55, [a], size=13, bold=True, color=INK)
        textbox(s, 3.5, y + 0.15, 4.0, 0.55, [b], size=13, color=SLATE)
        textbox(s, 7.7, y + 0.15, 5.0, 0.55, [c], size=13, color=SLATE)
        y += 0.75
    textbox(
        s,
        0.5,
        6.7,
        12,
        0.4,
        ["Best first placement: CI/CD + pre-prod governance as the mandatory checkpoint."],
        size=13,
        bold=True,
        color=TEAL,
    )

    # 7 Stack view
    s = prs.slides.add_slide(blank(prs))
    add_rect(s, 0, 0, 13.333, 7.5, PAPER)
    title_bar(s, "Stack view — where GIE sits")
    layers = [
        ("Above", "Business / GRC", "Policy intent, frameworks, audit asks", False),
        ("GIE", "Orchestration", "Choose & generate the right guardrails per agent", True),
        ("Beside", "Build tools", "IDE, Git, CI, ticketing", False),
        ("Below", "Agent runtime", "LangGraph, CrewAI, MCP, tools", False),
        ("Below", "Models / cloud", "LLM providers, infra", False),
    ]
    x = 0.35
    for tag, title, body, accent in layers:
        card(s, x, 1.5, 2.45, 3.6, f"{tag} · {title}", [body], accent=accent)
        x += 2.55
    add_rect(s, 0.5, 5.5, 12.3, 1.3, NAVY)
    textbox(
        s,
        0.75,
        5.85,
        11.8,
        0.8,
        ["GIE is not “another agent.” It is the middle governance layer that decides which controls apply before agents go live."],
        size=15,
        bold=True,
        color=WHITE,
    )

    # 8 Outcomes
    s = prs.slides.add_slide(blank(prs))
    add_rect(s, 0, 0, 13.333, 7.5, PAPER)
    title_bar(s, "Business outcomes")
    card(s, 0.5, 1.4, 6.0, 2.4, "Faster safe adoption", ["Ship agents with matched controls instead of waiting on ad-hoc policy reviews."])
    card(s, 6.8, 1.4, 6.0, 2.4, "Lower residual risk", ["Tool abuse, leakage, weak identity, and over-autonomy scored before production."])
    card(s, 0.5, 4.1, 6.0, 2.4, "Audit-ready decisions", ["Every run leaves an execution trail and plain-English release posture."])
    card(s, 6.8, 4.1, 6.0, 2.4, "Consistent org standard", ["One orchestration path for Sec, Platform, and Compliance."])

    # 9 USP
    s = prs.slides.add_slide(blank(prs))
    add_rect(s, 0, 0, 13.333, 7.5, PAPER)
    title_bar(s, "USP for customer conversations")
    bullets(
        s,
        0.7,
        1.4,
        12,
        4.2,
        [
            "Context-aware policy selection — controls matched to real agent capabilities",
            "End-to-end orchestration — discover → risk → recommend → generate → validate → explain",
            "Deployable artifacts — not only findings; policies ready for enforcement paths",
            "Auditable by design — deterministic scoring + execution evidence",
            "Agentic-native — tools, MCP, autonomy, multi-agent patterns in scope",
        ],
        size=18,
    )
    add_rect(s, 0.5, 5.8, 12.3, 1.1, NAVY)
    textbox(
        s,
        0.75,
        6.1,
        11.8,
        0.7,
        ['Pitch: “One orchestration layer so every AI agent gets the right guardrails before it ships.”'],
        size=15,
        bold=True,
        color=WHITE,
    )

    # 10 Is / Is not
    s = prs.slides.add_slide(blank(prs))
    add_rect(s, 0, 0, 13.333, 7.5, PAPER)
    title_bar(s, "What it is / is not")
    card(
        s,
        0.5,
        1.4,
        6.0,
        5.0,
        "Is",
        [
            "• Governance / orchestration middle layer",
            "• Release checkpoint for agentic apps",
            "• Selector + generator of guardrails/policies",
            "• Evidence plane for leadership & audit",
        ],
        accent=True,
    )
    card(
        s,
        6.8,
        1.4,
        6.0,
        5.0,
        "Is not (today)",
        [
            "• Inline runtime proxy for every user↔agent call",
            "• Replacement for SIEM / CSPM / classic GRC",
            "• A single chatbot “doing security”",
            "• Deep live PII redaction engine (signals today)",
        ],
    )

    # 11 Adoption
    s = prs.slides.add_slide(blank(prs))
    add_rect(s, 0, 0, 13.333, 7.5, PAPER)
    title_bar(s, "Recommended leverage path")
    card(s, 0.5, 1.5, 4.0, 3.8, "Phase 1", ["CI + pre-prod gate", "", "Mandatory analyze on agent repos; hold / conditional / clear."], accent=True)
    card(s, 4.7, 1.5, 4.0, 3.8, "Phase 2", ["Policy pack handoff", "", "Push generated guardrails into runtime firewalls / platform policy stores."])
    card(s, 8.9, 1.5, 4.0, 3.8, "Phase 3", ["Org standard", "", "All new agents register through GIE; GRC consumes evidence automatically."])
    textbox(
        s,
        0.5,
        5.7,
        12,
        1.0,
        ["Start where decisions already happen: release. Expand to runtime enforcement once selection quality is trusted."],
        size=15,
        color=SLATE,
    )

    # 12 Q&A + ask
    s = prs.slides.add_slide(blank(prs))
    add_rect(s, 0, 0, 13.333, 7.5, PAPER)
    title_bar(s, "Ask of this review")
    card(
        s,
        0.5,
        1.4,
        6.0,
        3.2,
        "Decision needed",
        ["Confirm GIE as the organization’s guardrail/policy orchestration layer for agentic systems."],
        accent=True,
    )
    card(
        s,
        6.8,
        1.4,
        6.0,
        3.2,
        "Next 30 days",
        [
            "• Pilot 2–3 priority agent apps in CI gate",
            "• Define release postures with Security",
            "• Map generated policies to current enforcement stack",
        ],
    )
    add_rect(s, 0.5, 5.0, 12.3, 1.7, NAVY)
    textbox(
        s,
        0.75,
        5.35,
        11.8,
        1.2,
        [
            "Management takeaway: leverage GIE first as the release orchestration layer;",
            "then connect it to runtime enforcement and GRC evidence.",
        ],
        size=16,
        bold=True,
        color=WHITE,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
