#!/usr/bin/env python3
"""Generate a viewer-compatible GIE Technical Deep-Dive PPTX.

Avoids theme style overrides that make some viewers (Google Slides,
LibreOffice, web previews) render slides as blank/white.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

INK = RGBColor(0x0A, 0x16, 0x28)
NAVY = RGBColor(0x10, 0x2A, 0x43)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
TEAL_BRIGHT = RGBColor(0x14, 0xB8, 0xA6)
MIST = RGBColor(0x9F, 0xB3, 0xC8)
FOG = RGBColor(0xD9, 0xE2, 0xEC)
PAPER = RGBColor(0xF0, 0xF4, 0xF8)
WHITE = RGBColor(0xFA, 0xFB, 0xFC)
SLATE = RGBColor(0x48, 0x65, 0x81)

OUT = Path(__file__).resolve().parents[1] / "presentations" / "GIE_Technical_Deep_Dive.pptx"
NSMAP = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def strip_style(shape) -> None:
    """Remove theme style refs that break some PPTX viewers."""
    sp = shape._element
    for style in sp.findall(qn("p:style")):
        sp.remove(style)


def set_solid(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    strip_style(shape)


def add_rect(slide, left, top, width, height, color: RGBColor):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    set_solid(sh, color)
    return sh


def add_round(slide, left, top, width, height, color: RGBColor):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    set_solid(sh, color)
    return sh


def textbox(slide, left, top, width, height, lines, *, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Calibri"
        p.space_after = Pt(4)
    return box


def bullets(slide, left, top, width, height, items, *, size=16, color=INK):
    lines = [f"•  {item}" for item in items]
    return textbox(slide, left, top, width, height, lines, size=size, color=color)


def footer(slide, page: int, total: int):
    textbox(
        slide,
        0.5,
        7.1,
        10,
        0.3,
        ["GIE Technical Deep-Dive  ·  Confidential"],
        size=10,
        color=MIST,
    )
    textbox(
        slide,
        11.4,
        7.1,
        1.4,
        0.3,
        [f"{page} / {total}"],
        size=10,
        color=MIST,
        align=PP_ALIGN.RIGHT,
    )


def title_bar(slide, title: str):
    add_rect(slide, 0, 0, 13.333, 0.95, NAVY)
    add_rect(slide, 0, 0.95, 13.333, 0.07, TEAL_BRIGHT)
    textbox(slide, 0.5, 0.25, 12.3, 0.55, [title], size=26, bold=True, color=WHITE)


def card(slide, left, top, width, height, title, body, *, fill=WHITE):
    add_round(slide, left, top, width, height, fill)
    # border via thin outer rect is heavy; keep simple fill
    textbox(slide, left + 0.18, top + 0.14, width - 0.36, 0.35, [title], size=14, bold=True, color=TEAL)
    textbox(slide, left + 0.18, top + 0.5, width - 0.36, height - 0.65, body, size=13, color=SLATE)


def add_table(slide, left, top, width, rows, col_w):
    nrows, ncols = len(rows), len(rows[0])
    table_shape = slide.shapes.add_table(
        nrows, ncols, Inches(left), Inches(top), Inches(width), Inches(0.38 * nrows)
    )
    table = table_shape.table
    for i, w in enumerate(col_w):
        table.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(11)
                    run.font.name = "Calibri"
                    run.font.bold = r == 0
                    run.font.color.rgb = WHITE if r == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (PAPER if r % 2 else WHITE)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides = []

    def slide(bg=PAPER):
        s = prs.slides.add_slide(blank)
        add_rect(s, 0, 0, 13.333, 7.5, bg)
        slides.append(s)
        return s

    # 1 Title
    s = slide(NAVY)
    textbox(s, 0.7, 2.1, 12, 1.0, ["GIE"], size=72, bold=True, color=TEAL_BRIGHT)
    textbox(s, 0.7, 3.3, 12, 0.5, ["Guardrails Intelligence Engine"], size=28, bold=True, color=WHITE)
    textbox(
        s,
        0.7,
        4.1,
        12,
        1.2,
        [
            "Technical Deep-Dive — Architecture, Agents, Pipeline, Data Plane & Deployment",
            "Discussion & demonstration deck  ·  Multi-agent AI governance platform",
        ],
        size=16,
        color=MIST,
    )

    # 2 Agenda
    s = slide()
    title_bar(s, "Agenda")
    bullets(
        s,
        0.6,
        1.3,
        5.8,
        5.4,
        [
            "Why GIE — problem & value",
            "Platform architecture overview",
            "Hexagonal / Clean Architecture pattern",
            "The 12-agent mesh (ports & schemas)",
            "Orchestrator pipeline & execution graph",
            "Contracts, envelopes & versioning",
            "Data plane (Postgres, Redis, Neo4j, Qdrant, Kafka)",
        ],
        size=16,
    )
    bullets(
        s,
        7.0,
        1.3,
        5.8,
        5.4,
        [
            "Security, authn/z & approval gates",
            "Observability & operations",
            "Local Docker validation path",
            "Azure full-capability sandbox",
            "Live demo walkthrough",
            "Integration surfaces (IDE, CI, APIs)",
            "Discussion topics & next steps",
        ],
        size=16,
    )

    # 3 Why
    s = slide()
    title_bar(s, "Why GIE exists")
    card(
        s,
        0.5,
        1.3,
        4.0,
        5.2,
        "The gap",
        [
            "AI apps ship with prompts, tools,",
            "MCP servers, model calls, and cloud",
            "bindings — often invisible to risk.",
            "",
            "Point scanners miss topology.",
            "Spreadsheets miss evidence.",
            "Chatbots miss audit trails.",
        ],
    )
    card(
        s,
        4.7,
        1.3,
        4.0,
        5.2,
        "What GIE does",
        [
            "Scan → normalize → reason →",
            "recommend → validate → explain.",
            "",
            "One Orchestrator entrypoint.",
            "Specialized agents own one job.",
            "Human gates when stakes rise.",
            "Traces you can defend in audit.",
        ],
    )
    card(
        s,
        8.9,
        1.3,
        4.0,
        5.2,
        "Outcomes",
        [
            "Shared Context Model (gie.context.v1)",
            "Risk + compliance posture",
            "Deployment-ready policies",
            "Audience-aware explanations",
            "Learning loop with approval",
            "Enterprise connectors / IDE",
        ],
    )

    # 4 Architecture
    s = slide()
    title_bar(s, "Platform architecture")
    textbox(
        s,
        0.5,
        1.2,
        12.3,
        0.4,
        ["Clients → Orchestrator (:8091) → peer agents (:8080–:8090) → shared data plane"],
        size=15,
        color=SLATE,
    )
    layers = [
        ("Clients", ["Dashboard / CLI", "IDE extensions", "CI / GitHub Action", "MCP tools"]),
        ("Control", ["Orchestrator", "Analyze / Approve", "Trace / Graph", "Status / Cache"]),
        ("Agents", ["Context ∥ Knowledge", "Risk ∥ Compliance", "Policy → Rec → Gen", "Validate / Explain / Learn"]),
        ("Data", ["PostgreSQL ×12 DBs", "Redis indexes 0–11", "Neo4j + Qdrant", "Kafka (KRaft)"]),
    ]
    for i, (t, body) in enumerate(layers):
        card(s, 0.5 + i * 3.15, 1.9, 3.0, 4.5, t, body)

    # 5 Hexagonal
    s = slide()
    title_bar(s, "Design pattern — Hexagonal / Clean Architecture")
    bullets(
        s,
        0.5,
        1.25,
        12.3,
        5.5,
        [
            "Every agent: Adapters → Application → Domain ← Ports ← Driven adapters",
            "Domain free of FastAPI, SQLAlchemy, Neo4j, Kafka — swapped via ports",
            "DDD: ContextModel aggregate; Scan process manager; Knowledge versioned graph",
            "CQRS: commands write Postgres + outbox; queries hit Postgres/Redis",
            "Event-driven: outbox → Kafka for cross-agent integration",
            "Shared packages: gie-contracts · gie-observability · gie-security",
            "Tests: unit (domain) · integration (API) · contract (OpenAPI) · smoke (compose/Azure)",
        ],
        size=16,
    )

    # 6 Agent table
    s = slide()
    title_bar(s, "Agent mesh — ports & schemas")
    add_table(
        s,
        0.35,
        1.2,
        12.6,
        [
            ["Agent", "Port", "Schema", "Primary responsibility"],
            ["Context Intelligence", "8080", "gie.context.v1", "Scan AI apps → Context Model + evidence"],
            ["Knowledge Intelligence", "8081", "gie.knowledge.v1", "Versioned knowledge graph + retrieval"],
            ["Policy Intelligence", "8082", "gie.policy.v1", "Guardrail policy selection"],
            ["Risk Intelligence", "8083", "gie.risk.v1", "AI risk posture + remediations"],
            ["Compliance Intelligence", "8084", "gie.compliance.v1", "Framework gaps + evidence"],
            ["Recommendation", "8085", "gie.recommendation.v1", "Prioritized actions"],
            ["Policy Generator", "8086", "gie.policygen.v1", "Deployment-ready policies"],
            ["Explainability", "8087", "gie.explainability.v1", "Multi-audience narratives"],
            ["Validation", "8088", "gie.validation.v1", "Pre-deploy validate / simulate"],
            ["Learning", "8089", "gie.learning.v1", "Feedback → knowledge proposals"],
            ["Integration", "8090", "gie.integration.v1", "Enterprise connectors + IDE"],
            ["Orchestrator", "8091", "gie.orchestrator.v1", "Unified analyze + gates + traces"],
        ],
        [3.1, 1.0, 2.5, 6.0],
    )

    # 7 Orchestrator
    s = slide()
    title_bar(s, "Orchestrator — unified analyze pipeline")
    bullets(
        s,
        0.5,
        1.2,
        6.0,
        5.5,
        [
            "POST /analyze — single entry for the mesh",
            "Modes: sync · async · streaming · batch",
            "Topological waves with parallel groups",
            "Risk ∥ Compliance fan-out where safe",
            "Retries with backoff; optional cache",
            "SIMULATE_AGENTS=true → SimulatedInvoker",
            "SIMULATE_AGENTS=false → HttpInvoker",
            "GET /status · /execution/{id} · /trace · /graph",
            "POST /approve — human gate for waiting steps",
        ],
        size=15,
    )
    card(
        s,
        6.8,
        1.2,
        5.9,
        5.5,
        "AnalyzeRequest highlights",
        [
            "tenant_id (required)",
            "source: { type, path | repo | zip }",
            "mode: sync | async | streaming | batch",
            "workflow_id: gie.analyze.default",
            "options: { parallel, audience, … }",
            "require_human_approval / approval_gates",
            "cache · timeout_ms · agent_versions",
            "",
            "Response: ObservabilityEnvelope",
            "  data = ExecutionRecord",
            "  meta = trace_id, confidence, ms",
        ],
    )

    # 8 Graph
    s = slide()
    title_bar(s, "Default execution graph (conceptual)")
    steps = [
        ("01 Context", "Materialize source, run detectors, persist Context Model"),
        ("02 Knowledge", "Retrieve frameworks / attacks / guardrails for context"),
        ("03 Risk ∥ Compliance", "Parallel posture + regulatory gap analysis"),
        ("04 Policy", "Select applicable guardrail policies"),
        ("05 Recommend → Generate", "Prioritize actions, emit deployable policies"),
        ("06 Validate", "Simulate / pre-deploy checks"),
        ("07 Explain", "Audience-aware narrative (eng / risk / exec)"),
        ("08 Gate / Learn", "Approve publishes; feed learning loop"),
    ]
    for i, (name, desc) in enumerate(steps):
        y = 1.2 + (i % 4) * 1.35
        x = 0.5 if i < 4 else 6.9
        card(s, x, y, 5.9, 1.2, name, [desc])

    # 9 Context
    s = slide()
    title_bar(s, "Context Intelligence — deep dive")
    bullets(
        s,
        0.5,
        1.25,
        12.3,
        5.5,
        [
            "Input: folder / repo / zip → isolated Scan workspace",
            "Detectors → DetectionFinding → Merger → ContextModel (gie.context.v1)",
            "Captures: frameworks, models, prompts, MCP servers, secrets (fingerprinted), topology",
            "Persistence: PostgreSQL · Neo4j · Qdrant · Celery worker on Redis",
            "APIs: REST + CLI (gie-context) + GitHub webhook + MCP",
            "Ready: database + redis required; neo4j/qdrant reported best-effort",
            "Events: ContextModelUpdated for Risk / Policy / Compliance",
        ],
        size=16,
    )

    # 10 Knowledge
    s = slide()
    title_bar(s, "Knowledge Intelligence — platform brain")
    bullets(
        s,
        0.5,
        1.25,
        12.3,
        5.5,
        [
            "Versioned knowledge graph: OWASP LLM, MITRE ATLAS, NIST AI RMF, controls, attacks, guardrails",
            "Explainable semantic retrieval over graph + embeddings",
            "Seed loader bootstraps baseline graph on start (sandbox/local)",
            "Production path: Postgres + Neo4j + Qdrant + Redis + Kafka",
            "Graceful fallback to in-memory if data plane fails (logged)",
            "Worker queue knowledge for reindex / expansion jobs",
            "Contract: gie.knowledge.v1",
        ],
        size=16,
    )

    # 11 Mid agents
    s = slide()
    title_bar(s, "Risk · Compliance · Policy · Rec · Generator")
    card(s, 0.4, 1.2, 4.1, 2.5, "Risk (:8083)", ["AI risk posture scoring", "Remediation suggestions", "gie.risk.v1"])
    card(s, 4.6, 1.2, 4.1, 2.5, "Compliance (:8084)", ["Catalog-driven frameworks", "Gaps + evidence", "gie.compliance.v1"])
    card(s, 8.8, 1.2, 4.1, 2.5, "Policy (:8082)", ["Guardrail selection", "Runtime policy packs", "gie.policy.v1"])
    card(s, 0.4, 4.0, 4.1, 2.5, "Recommendation (:8085)", ["Prioritized backlog", "Impact × effort", "gie.recommendation.v1"])
    card(s, 4.6, 4.0, 4.1, 2.5, "Policy Generator (:8086)", ["Recs → deployable artifacts", "gie.policygen.v1"])
    card(s, 8.8, 4.0, 4.1, 2.5, "Parallelism", ["Risk ∥ Compliance in default", "workflow parallel_group", "max_parallel_steps"])

    # 12 Later agents
    s = slide()
    title_bar(s, "Explain · Validate · Learn · Integrate")
    card(s, 0.4, 1.2, 6.1, 2.6, "Explainability (:8087)", ["Multi-audience narratives: engineering · risk · executive", "Ties decisions to evidence + confidence", "gie.explainability.v1"])
    card(s, 6.7, 1.2, 6.1, 2.6, "Validation (:8088)", ["Pre-deploy validate + simulate", "Catch policy regressions before promotion", "gie.validation.v1"])
    card(s, 0.4, 4.1, 6.1, 2.6, "Learning (:8089)", ["Feedback → knowledge proposals", "ALLOW_AUTO_PUBLISH=false by default", "Human approval before publish", "gie.learning.v1"])
    card(s, 6.7, 4.1, 6.1, 2.6, "Integration (:8090)", ["Enterprise connectors", "extensions/: VS Code · Cursor · JetBrains · GitHub Action", "gie.integration.v1"])

    # 13 Contracts
    s = slide()
    title_bar(s, "Contracts & API envelopes")
    bullets(
        s,
        0.5,
        1.25,
        12.3,
        5.5,
        [
            "packages/gie-contracts — Pydantic models, versioned schema strings (gie.*.v1)",
            "REST aliases: /v1/... and unversioned shortcuts (e.g. POST /analyze)",
            "ObservabilityEnvelope { data, meta } — trace_id, request_id, confidence, agent_version",
            "ExecutionRecord: execution_id, steps[], status, result, reasoning_path, cache_hits",
            "StepStatus: pending · running · succeeded · failed · skipped · retrying · cached · waiting_approval",
            "Breaking changes → new schema version; agent_versions map allows routing overrides",
            "OpenAPI at each agent /docs; contract tests assert path presence",
        ],
        size=15,
    )

    # 14 Data plane
    s = slide()
    title_bar(s, "Data plane")
    add_table(
        s,
        0.35,
        1.2,
        12.6,
        [
            ["Store", "Role", "Local", "Azure sandbox"],
            ["PostgreSQL 16", "Primary write model / outbox", "compose service", "Flexible Server · 12 DBs"],
            ["Redis 7", "Cache + Celery broker", "compose", "Azure Cache Basic C0"],
            ["Neo4j 5", "Graph projection / knowledge", "compose", "In-cluster Deployment"],
            ["Qdrant 1.12", "Evidence / embeddings", "compose", "In-cluster Deployment"],
            ["Kafka 3.7.1", "Domain events (KRaft)", "apache/kafka", "In-cluster Deployment"],
        ],
        [2.2, 3.5, 2.8, 4.1],
    )
    bullets(
        s,
        0.5,
        4.5,
        12.3,
        2.2,
        [
            "Logical DBs: gie_context … gie_orchestrator (init script on first Postgres boot)",
            "Redis DB indexes 0–11 isolate agent namespaces on one Redis instance",
            "Orchestrator API currently boots in-memory repos; platform DBs serve Context/Knowledge today",
        ],
        size=14,
    )

    # 15 Security
    s = slide()
    title_bar(s, "Security & approval gates")
    bullets(
        s,
        0.5,
        1.25,
        12.3,
        5.5,
        [
            "Auth: JWT (HS256) and/or API keys (gie_…); REQUIRE_AUTH toggles enforcement",
            "Sandbox: REQUIRE_AUTH=false with x-tenant-id for tenancy",
            "RBAC: OrchestratorPermission checked in deps",
            "Secrets fingerprinted in Context — raw secrets not stored in models",
            "Approval: require_human_approval → WAITING_APPROVAL → POST /approve",
            "Learning publish blocked unless approved (ALLOW_AUTO_PUBLISH=false)",
            "CORS permissive when require_auth=false or gie_env in {local,test,docker,dev,azure}",
            "Pilot hardening: Entra ID, Key Vault CSI, private AKS, WAF (Phase 4)",
        ],
        size=15,
    )

    # 16 Observability
    s = slide()
    title_bar(s, "Observability & operations")
    card(s, 0.5, 1.2, 4.0, 5.3, "Signals", ["structlog JSON stdout", "OpenTelemetry → OTLP", "prometheus /metrics", "/healthz · /ready", "Execution traces per run", "Mermaid graph endpoint"])
    card(s, 4.7, 1.2, 4.0, 5.3, "Ops tips", ["scripts/local-smoke.sh", "Azure: 05-smoke.sh", "Dashboard → API base URL", "kubectl logs -n gie", "Check SIMULATE_AGENTS", "Peer /healthz from orch"])
    card(s, 8.9, 1.2, 4.0, 5.3, "Quality", ["pytest unit/integration", "OpenAPI contract tests", "Helm lint charts", "ruff + mypy", "Compose build matrix", "Azure teardown when idle"])

    # 17 Local
    s = slide()
    title_bar(s, "Local Docker — full mesh validation")
    bullets(
        s,
        0.5,
        1.25,
        12.3,
        5.5,
        [
            "./scripts/local-up.sh   then   ./scripts/local-smoke.sh",
            "SIMULATE_AGENTS=false — HttpAgentInvoker hits real peers",
            "Web UI :5173 · Orchestrator :8091 · Agents :8080–:8090",
            "docker-compose.host.yml when Docker bridge/iptables is broken",
            "LOCAL_NET_MODE=bridge on normal Docker Desktop / Linux hosts",
            "Validated: all peer healthz · status healthy · analyze live steps",
            "Context /ready: database + redis + neo4j + qdrant true",
            "Docs: deploy/docker/README.md",
        ],
        size=15,
    )

    # 18 Azure
    s = slide()
    title_bar(s, "Azure full-capability sandbox")
    bullets(
        s,
        0.5,
        1.2,
        6.2,
        5.5,
        [
            "SANDBOX_MODE=full (default)",
            "AKS 3× Standard_D4s_v5",
            "ACR + Postgres (12 DBs) + Redis",
            "In-cluster Neo4j / Qdrant / Kafka",
            "All 12 agents + workers + web LB",
            "Orchestrator SIMULATE_AGENTS=false",
            "Cost ~$400–700+/mo — tear down idle",
            "minimal mode: Orchestrator-only simulate",
        ],
        size=15,
    )
    card(
        s,
        6.9,
        1.2,
        5.9,
        5.5,
        "Deploy sequence",
        [
            "00-prereqs.sh",
            "01-infra.sh",
            "02-build-push.sh   (12 images + web)",
            "03-deploy-data.sh  (Neo4j/Qdrant/Kafka)",
            "04-deploy-app.sh   (Helm + workers + web)",
            "05-smoke.sh",
            "",
            "Outputs: .last-url.txt · .last-web-url.txt",
            "See deploy/azure/README.md",
        ],
    )

    # 19 Demo
    s = slide()
    title_bar(s, "Live demo script")
    bullets(
        s,
        0.5,
        1.25,
        12.3,
        5.5,
        [
            "Open landing http://127.0.0.1:5173 — brand-first hero, pipeline, agents",
            "Open dashboard — API base http://127.0.0.1:8091 — Ping / Refresh status",
            "Show GET / and /docs on Orchestrator — contract surface",
            "Run Analyze (sync): tenant local-demo, source /demo — watch steps succeed",
            "Fetch execution by id — show ExecutionRecord + step outputs",
            "Optional: curl Context /ready — prove data plane green",
            "Flip narrative: Azure path uses same contracts, different topology",
            "Call out approval gate + Learning publish for regulated workflows",
        ],
        size=15,
    )

    # 20 Integrations
    s = slide()
    title_bar(s, "Integration surfaces")
    card(s, 0.4, 1.2, 4.1, 5.3, "APIs", ["REST (FastAPI)", "OpenAPI /docs", "MCP adapters", "CLI per agent", "Webhooks (GitHub)"])
    card(s, 4.6, 1.2, 4.1, 5.3, "Extensions", ["VS Code", "Cursor", "JetBrains", "GitHub Action", "under extensions/"])
    card(s, 8.8, 1.2, 4.1, 5.3, "Deploy targets", ["docker compose", "Helm charts ×12", "Azure Bicep sandbox", "Terraform (AWS-oriented)", "CI: GitHub Actions"])

    # 21 Discussion
    s = slide()
    title_bar(s, "Discussion topics")
    bullets(
        s,
        0.5,
        1.25,
        12.3,
        5.5,
        [
            "Where should human approval sit in your SDLC (PR, staging, prod promote)?",
            "Which agents are must-have for pilot vs phase-2 (full mesh)?",
            "Tenant model: single enterprise tenant vs business-unit isolation?",
            "Data residency: keep Neo4j/Qdrant in-cluster vs managed services?",
            "Auth: API keys for CI vs Entra ID for humans?",
            "Success metrics: MTTR for findings · % policies auto-validated · audit time",
            "Open: wire Orchestrator persistence off in-memory → Postgres/Redis",
            "Open: deepen HttpAgentInvoker beyond /healthz into full peer API payloads",
        ],
        size=15,
    )

    # 22 Next
    s = slide()
    title_bar(s, "Recommended next steps")
    card(s, 0.5, 1.3, 4.0, 5.2, "Now", ["Keep local compose green", "Walk dashboard analyze", "Pick pilot use-case", "Define approval owners"])
    card(s, 4.7, 1.3, 4.0, 5.2, "Next", ["Azure SANDBOX_MODE=full", "Connect real repo source", "Enable REQUIRE_AUTH", "CI Action in one repo"])
    card(s, 8.9, 1.3, 4.0, 5.2, "Later", ["Private AKS + WAF", "Key Vault CSI", "Entra ID", "HA data plane", "Full invoker APIs"])

    # 23 Close
    s = slide(NAVY)
    textbox(s, 0.7, 2.5, 12, 0.7, ["Govern AI decisions before they ship."], size=32, bold=True, color=TEAL_BRIGHT)
    textbox(
        s,
        0.7,
        3.5,
        12,
        1.5,
        [
            "GIE — Guardrails Intelligence Engine",
            "Questions & discussion",
            "docs/  ·  deploy/azure/  ·  deploy/docker/  ·  web/",
        ],
        size=16,
        color=MIST,
    )

    total = len(slides)
    for i, s in enumerate(slides, start=1):
        if i not in (1, total):
            footer(s, i, total)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)

    # Verify
    check = Presentation(str(OUT))
    assert len(check.slides) == total
    s0 = check.slides[0]
    assert any(sh.has_text_frame and "GIE" in sh.text_frame.text for sh in s0.shapes)
    # ensure no p:style left on first shape
    xml = s0.shapes[0]._element.xml
    assert "<p:style>" not in xml, "theme style still present — viewers may blank out"
    print(f"Wrote {OUT} ({total} slides) — verified text + no theme styles")


if __name__ == "__main__":
    build()
